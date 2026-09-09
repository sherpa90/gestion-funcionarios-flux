import os
import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'core.settings')
django.setup()

from datetime import date, time, timedelta
from django.utils.timezone import now
from users.models import CustomUser
from asistencia.models import RegistroAsistencia, HorarioFuncionario, DiaHorario, DiaFestivo, AnoEscolar
from permisos.models import SolicitudPermiso
from licencias.models import LicenciaMedica
from reportes.calc_utils import calcular_metricas_cla_en_lote
from reportes.views import ExportarCLAExcelView, ExportarCLAPDFView
from django.test import RequestFactory
import openpyxl
import io

def run_tests():
    print("=== INICIANDO PRUEBAS DE INFORME CLA ===")

    # 1. Crear o recuperar año escolar 2026
    ano_escolar, _ = AnoEscolar.objects.get_or_create(
        ano=2026,
        defaults={
            'sem1_inicio': date(2026, 3, 1),
            'sem1_fin': date(2026, 7, 10),
            'sem2_inicio': date(2026, 7, 27),
            'sem2_fin': date(2026, 12, 15),
        }
    )

    # Limpiar usuario previo si existe
    CustomUser.objects.filter(username='test_docente_cla').delete()

    # 2. Crear usuario de prueba docente
    user_test = CustomUser.objects.create_user(
        username='test_docente_cla',
        first_name='Juan',
        last_name='Pérez CLA',
        email='juan.cla@test.com',
        password='testpass123',
        run='11.111.111-1',
        tipo_funcionario='DOCENTE',
        role='DOCENTE',
    )
    user_test.date_joined = now().replace(year=2025, month=1, day=1)
    user_test.save()

    # Crear horario lunes a viernes
    horario = HorarioFuncionario.objects.create(
        funcionario=user_test,
        activo=True,
        nombre='Horario Test Docente'
    )
    for d_sem in range(5):  # 0-4 = lunes-viernes
        DiaHorario.objects.create(
            horario=horario,
            dia_semana=d_sem,
            semana_tipo=None,
            hora_entrada=time(8, 0),
            hora_salida=time(16, 0),
            activo=True
        )

    # Marzo 2026: 2026-03-02 es Lunes

    # Caso 1: 2026-03-02 (Lunes) -> Falta entrada (tiene solo salida) -> SIN ENTRADA
    RegistroAsistencia.objects.create(
        funcionario=user_test,
        fecha=date(2026, 3, 2),
        hora_entrada_real=None,
        hora_salida_real=time(16, 0),
        estado='SIN_MARCACION_ENTRADA'
    )

    # Caso 2: 2026-03-03 (Martes) -> Falta salida (tiene solo entrada) -> SIN SALIDA
    RegistroAsistencia.objects.create(
        funcionario=user_test,
        fecha=date(2026, 3, 3),
        hora_entrada_real=time(8, 0),
        hora_salida_real=None,
        estado='PUNTUAL'
    )

    # Caso 3: 2026-03-04 (Miércoles) -> Sin entrada ni salida -> SIN AMBAS / INASISTENCIA
    RegistroAsistencia.objects.create(
        funcionario=user_test,
        fecha=date(2026, 3, 4),
        hora_entrada_real=None,
        hora_salida_real=None,
        estado='AUSENTE'
    )

    # Caso 4: 2026-03-05 (Jueves) -> Licencia médica (NO debe contar nada)
    LicenciaMedica.objects.create(
        usuario=user_test,
        tipo='LICENCIA',
        fecha_inicio=date(2026, 3, 5),
        dias=1,
        created_by=user_test
    )

    # Caso 5: 2026-03-06 (Viernes) -> Día administrativo completo (NO debe contar nada)
    SolicitudPermiso.objects.create(
        usuario=user_test,
        fecha_inicio=date(2026, 3, 6),
        fecha_termino=date(2026, 3, 6),
        dias_solicitados=1.0,
        jornada='FD',
        estado='APROBADO',
        observacion='Trámites personales'
    )

    # Caso 6: 2026-03-09 (Lunes) -> Medio día AM aprobado + sin entrada + con salida.
    # Permiso AM: falta de entrada AUTORIZADA → NO debe contar como "Sin Entrada"
    SolicitudPermiso.objects.create(
        usuario=user_test,
        fecha_inicio=date(2026, 3, 9),
        fecha_termino=date(2026, 3, 9),
        dias_solicitados=0.5,
        jornada='AM',
        estado='APROBADO',
        observacion='Medio día AM'
    )
    RegistroAsistencia.objects.create(
        funcionario=user_test,
        fecha=date(2026, 3, 9),
        hora_entrada_real=None,
        hora_salida_real=time(16, 0),
        estado='MEDIO_DIA'
    )

    # Caso 7: 2026-03-10 (Martes) -> Justificación manual. Sin entrada pero justificado. NO debe contar.
    RegistroAsistencia.objects.create(
        funcionario=user_test,
        fecha=date(2026, 3, 10),
        hora_entrada_real=None,
        hora_salida_real=time(16, 0),
        estado='JUSTIFICADO',
        justificacion_manual='Olvidó tarjeta pero dio aviso a secretaría',
        justificado_por=user_test
    )

    print("Datos de prueba creados exitosamente.")

    # Ejecutar cálculo
    metricas = calcular_metricas_cla_en_lote([user_test], 2026, 3)
    assert len(metricas) == 1, f"Se esperaba 1 funcionario en métricas, se obtuvo {len(metricas)}"
    m = metricas[0]
    print(f"Métricas calculadas:")
    print(f"  atrasos_formato: {m['atrasos_formato']}")
    print(f"  inasistencias:   {m['inasistencias']}")
    print(f"  sin_entrada:     {m['sin_entrada']}")
    print(f"  sin_salida:      {m['sin_salida']}")
    print(f"  sin_ambas:       {m['sin_ambas']}")

    # Verificaciones
    assert m['sin_entrada'] == 1, f"Esperado sin_entrada=1, obtenido={m['sin_entrada']}"
    print("✓ sin_entrada=1 (día 03-02, falta entrada sin permiso AM)")

    assert m['sin_salida'] == 1, f"Esperado sin_salida=1, obtenido={m['sin_salida']}"
    print("✓ sin_salida=1 (día 03-03, falta salida sin permiso PM)")

    assert m['sin_ambas'] == 1, f"Esperado sin_ambas=1, obtenido={m['sin_ambas']}"
    assert m['inasistencias'] == 1, f"Esperado inasistencias=1, obtenido={m['inasistencias']}"
    print("✓ sin_ambas=inasistencias=1 (día 03-04, sin entrada ni salida)")
    print("✓ Licencia médica 03-05 excluida correctamente")
    print("✓ Día administrativo 03-06 excluido correctamente")
    print("✓ Permiso AM 03-09 no contabilizó como sin_entrada")
    print("✓ Justificación manual 03-10 excluida correctamente")

    # Probar generación Excel
    rf = RequestFactory()
    req = rf.get('/reportes/cla-excel/?tipo=docente&year=2026&mes=3')
    req.user = user_test
    req.user.role = 'ADMIN'

    view_excel = ExportarCLAExcelView()
    view_excel.request = req
    resp_excel = view_excel.get(req)
    assert resp_excel.status_code == 200, f"Error en Excel: status {resp_excel.status_code}"
    assert 'spreadsheetml' in resp_excel['Content-Type']

    wb = openpyxl.load_workbook(io.BytesIO(resp_excel.content))
    ws = wb.active
    headers = [ws.cell(row=9, column=c).value for c in range(1, 8)]
    expected_headers = ['Nombre y Apellidos', 'RUN', 'Atrasos', 'Inasistencias', 'Sin Entrada', 'Sin Salida', 'Sin Ambas']
    assert headers == expected_headers, f"Encabezados Excel erróneos: {headers}"
    print(f"✓ Encabezados Excel: {headers}")

    fila = [ws.cell(row=10, column=c).value for c in range(1, 8)]
    print(f"✓ Fila datos Excel: {fila}")
    assert fila[4] == 1, f"Excel columna Sin Entrada debe ser 1, obtenido={fila[4]}"
    assert fila[5] == 1, f"Excel columna Sin Salida debe ser 1, obtenido={fila[5]}"
    assert fila[6] == 1, f"Excel columna Sin Ambas debe ser 1, obtenido={fila[6]}"
    print("✓ Valores de omisiones en Excel verificados")

    # Probar generación PDF
    req_pdf = rf.get('/reportes/cla-pdf/?tipo=docente&year=2026&mes=3')
    req_pdf.user = user_test
    req_pdf.user.role = 'ADMIN'

    view_pdf = ExportarCLAPDFView()
    view_pdf.request = req_pdf
    resp_pdf = view_pdf.get(req_pdf)
    assert resp_pdf.status_code == 200, f"Error en PDF: status {resp_pdf.status_code}"
    assert resp_pdf['Content-Type'] == 'application/pdf'
    assert resp_pdf.content.startswith(b'%PDF'), "Contenido no es PDF válido"
    print(f"✓ PDF generado exitosamente ({len(resp_pdf.content)} bytes)")

    # Limpiar
    CustomUser.objects.filter(username='test_docente_cla').delete()
    print("\n=== TODAS LAS PRUEBAS PASARON ===")

if __name__ == '__main__':
    run_tests()
